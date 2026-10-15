# -*- coding: utf-8 -*-
"""Export a native, editable Word file whose typography matches the PDF."""
from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "monograph"
SUBMIT_DIR = ROOT / "submit"
SUBMIT_PDF = SUBMIT_DIR / "20241334庄颂信息2402.pdf"
SUBMIT_DOCX = SUBMIT_DIR / "20241334庄颂信息2402.docx"
LOCAL_DOCX = BOOK / "ACTS-MCDT.docx"
PDF = BOOK / "ACTS-MCDT.pdf"
COVER = BOOK / "figures" / "cover.png"
BACK = BOOK / "figures" / "back.png"
BG = BOOK / "figures" / "bg.png"

A4_W = Mm(210)
A4_H = Mm(297)
BODY_FIG_MAX = Mm(148)
SONG = "宋体"
HEI = "黑体"
KAI = "楷体"
TIMES = "Times New Roman"
TOC_TAB_TWIPS = "9072"
BODY_LINE = 360
HEAD_LINE = 276
TOC_LINE = 280
HEADER_LEFT = "基于信息感知的智能蒸烤一体机"
HEADER_RIGHT = "庄颂  20241334  信息2402"

BODY_CHAPTERS = [
    "绪论",
    "文献基础与商品方案的信息边界",
    "感知、决策与执行的分层",
    "从全场物理到决策状态",
    "内部热阻的不可辨识性",
    "热信息探针与对偶决策",
    "因果路径与安全隔离",
    "厚度、档位与探针的对照结果",
    "网格收敛与冷点取法",
    "因子消融与边界扰动",
    "先验宽度、集合散布与时域",
    "冷点位置与过程尺度",
    "典型轨迹与作用机理",
    "数值实现与交互演示",
    "结论",
]
APPENDIX_CHAPTERS = [
    "符号、单位与日志字段",
    "一维工程核的离散格式",
    "种子 7 轨迹抽样",
    "用已定常数算出的尺度",
]
UNNUMBERED = {"摘要", "Abstract", "目录", "参考文献"}

PANDOC_PREAMBLE = r"""
\documentclass[12pt,a4paper]{book}
\usepackage{amsmath,amssymb,bm,graphicx,booktabs,longtable,hyperref,url,amsthm}
\newcommand{\R}{\mathbb{R}}
\newcommand{\E}{\mathbb{E}}
\newcommand{\Prob}{\mathbb{P}}
\newcommand{\cH}{\mathcal{H}}
\newcommand{\cF}{\mathcal{F}}
\newcommand{\cU}{\mathcal{U}}
\newcommand{\cX}{\mathcal{X}}
\newcommand{\cY}{\mathcal{Y}}
\newcommand{\cA}{\mathcal{A}}
\newcommand{\diag}{\mathrm{diag}}
\newcommand{\tr}{\mathrm{tr}}
\newcommand{\argmin}{\mathrm{arg\,min}}
\newcommand{\argmax}{\mathrm{arg\,max}}
\newcommand{\vect}[1]{\mathbf{#1}}
\newtheorem{definition}{定义}[chapter]
\newtheorem{proposition}{命题}[chapter]
\newtheorem{theorem}{定理}[chapter]
\newtheorem{lemma}{引理}[chapter]
\newtheorem{assumption}{假设}[chapter]
\begin{document}
""".strip()

INPUT_RE = re.compile(r"\\input\{([^}]+)\}")
TOKEN_RE = re.compile(
    r"\\chapter\*?\{|"
    r"\\section\*?\{|"
    r"\\subsection\*?\{|"
    r"\\appendix\b|"
    r"\\begin\{equation\*?\}|"
    r"\\end\{equation\*?\}|"
    r"\\begin\{align\*?\}|"
    r"\\end\{align\*?\}|"
    r"\\begin\{figure\}|"
    r"\\begin\{table\}|"
    r"\\begin\{longtable\}|"
    r"\\begin\{proposition\}|"
    r"\\begin\{definition\}|"
    r"\\begin\{theorem\}|"
    r"\\label\{([^}]+)\}"
)


def expand_inputs(tex: str, base: Path, skip: set[str] | None = None) -> str:
    skip = skip or set()

    def repl(match: re.Match[str]) -> str:
        name = match.group(1).strip()
        stem = Path(name).name
        if name in skip or stem in skip or stem.replace(".tex", "") in skip:
            return ""
        path = base / name
        if path.suffix == "":
            path = path.with_suffix(".tex")
        if not path.exists():
            return match.group(0)
        return expand_inputs(path.read_text(encoding="utf-8"), base, skip)

    return INPUT_RE.sub(repl, tex)


def strip_comments(tex: str) -> str:
    lines = []
    for line in tex.splitlines():
        stripped = line.lstrip()
        if stripped.startswith("%"):
            continue
        out = []
        i = 0
        while i < len(line):
            ch = line[i]
            if ch == "%" and (i == 0 or line[i - 1] != "\\"):
                break
            out.append(ch)
            i += 1
        lines.append("".join(out).rstrip())
    return "\n".join(lines)


def extract_cite_keys(tex: str) -> dict[str, int]:
    keys: dict[str, int] = {}
    for i, match in enumerate(re.finditer(r"\\bibitem\{([^}]+)\}", tex), start=1):
        keys[match.group(1)] = i
    return keys


def replace_cites(tex: str, keys: dict[str, int]) -> str:
    def repl(match: re.Match[str]) -> str:
        nums = []
        for key in match.group(1).split(","):
            key = key.strip()
            nums.append(str(keys.get(key, key)))
        return "[" + ",".join(nums) + "]"

    return re.sub(r"\\cite\{([^}]+)\}", repl, tex)


def convert_bibliography(tex: str) -> str:
    pattern = re.compile(
        r"\\begin\{thebibliography\}\{\d+\}(.*?)\\end\{thebibliography\}",
        re.S,
    )

    def repl(match: re.Match[str]) -> str:
        body = match.group(1)
        items = re.split(r"\\bibitem\{[^}]+\}", body)
        items = [item.strip() for item in items if item.strip()]
        lines = ["\\chapter*{参考文献}", "\\begin{enumerate}"]
        for item in items:
            item = item.replace("\\newblock", " ")
            item = item.replace("~", " ")
            item = re.sub(r"[ \t]+\n", "\n", item)
            item = re.sub(r"\s+", " ", item).strip()
            lines.append("\\item " + item)
        lines.append("\\end{enumerate}")
        return "\n".join(lines)

    return pattern.sub(repl, tex)


def convert_tabularx(tex: str) -> str:
    def repl(match: re.Match[str]) -> str:
        spec = match.group(1).replace("X", "p{5.4cm}")
        return r"\begin{tabular}{" + spec + "}"

    tex = re.sub(r"\\begin\{tabularx\}\{[^}]*\}\{([^}]*)\}", repl, tex)
    return tex.replace(r"\end{tabularx}", r"\end{tabular}")


def collect_labels(tex: str) -> tuple[dict[str, str], set[str]]:
    labels: dict[str, str] = {}
    eq_ids: set[str] = set()
    chapter = 0
    section = 0
    subsection = 0
    eq_n = 0
    fig_n = 0
    tab_n = 0
    prop_n = 0
    def_n = 0
    appendix = False
    app_i = 0
    in_eq = False
    pending: tuple[str, str] | None = None

    def chap_id() -> str:
        if appendix:
            return chr(ord("A") + app_i - 1)
        return str(chapter)

    for match in TOKEN_RE.finditer(tex):
        token = match.group(0)
        if token.startswith(r"\chapter*"):
            pending = None
            continue
        if token.startswith(r"\chapter"):
            if appendix:
                app_i += 1
            else:
                chapter += 1
            section = 0
            subsection = 0
            eq_n = fig_n = tab_n = prop_n = def_n = 0
            pending = ("ch", chap_id())
            continue
        if token == r"\appendix":
            appendix = True
            chapter = 0
            app_i = 0
            pending = None
            continue
        if token.startswith(r"\section"):
            section += 1
            subsection = 0
            pending = ("sec", f"{chap_id()}.{section}")
            continue
        if token.startswith(r"\subsection"):
            subsection += 1
            pending = ("sub", f"{chap_id()}.{section}.{subsection}")
            continue
        if token.startswith(r"\begin{equation") or token.startswith(r"\begin{align"):
            in_eq = True
            continue
        if token.startswith(r"\end{equation") or token.startswith(r"\end{align"):
            in_eq = False
            continue
        if token.startswith(r"\begin{figure}"):
            fig_n += 1
            pending = ("fig", f"{chap_id()}.{fig_n}")
            continue
        if token.startswith(r"\begin{table}") or token.startswith(r"\begin{longtable}"):
            tab_n += 1
            pending = ("tab", f"{chap_id()}.{tab_n}")
            continue
        if token.startswith(r"\begin{proposition}"):
            prop_n += 1
            pending = ("prop", f"{chap_id()}.{prop_n}")
            continue
        if token.startswith(r"\begin{definition}") or token.startswith(r"\begin{theorem}"):
            def_n += 1
            pending = ("def", f"{chap_id()}.{def_n}")
            continue
        label = match.group(1)
        if not label:
            continue
        if in_eq:
            eq_n += 1
            labels[label] = f"{chap_id()}.{eq_n}"
            eq_ids.add(label)
        elif pending is not None:
            labels[label] = pending[1]
    return labels, eq_ids


def replace_refs(tex: str, labels: dict[str, str], eq_ids: set[str]) -> str:
    def eqref(match: re.Match[str]) -> str:
        key = match.group(1)
        return "(" + labels.get(key, key) + ")"

    def ref(match: re.Match[str]) -> str:
        key = match.group(1)
        num = labels.get(key, key)
        if key in eq_ids:
            return num
        return num

    tex = re.sub(r"\\eqref\{([^}]+)\}", eqref, tex)
    tex = re.sub(r"\\ref\{([^}]+)\}", ref, tex)
    return tex


def strip_layout(tex: str) -> str:
    tex = re.sub(r"\\FullBleedPage\{[^}]*\}", "", tex)
    tex = re.sub(r"\\BodyBg(true|false)", "", tex)
    tex = re.sub(r"\\pagestyle\{[^}]*\}", "", tex)
    tex = re.sub(r"\\setcounter\{page\}\{[^}]*\}", "", tex)
    tex = re.sub(r"\\tableofcontents", "", tex)
    tex = re.sub(r"\\FloatBarrier", "", tex)
    tex = re.sub(r"\\clearpage", r"\\newpage", tex)
    tex = re.sub(r"\\begin\{proof\}", r"\n\\noindent\\textbf{证明.}\n", tex)
    tex = re.sub(r"\\end\{proof\}", "", tex)
    tex = tex.replace(r"\operatorname{ess\,inf}", r"\mathrm{ess\,inf}")
    tex = tex.replace(r"\operatorname*{arg\,min}", r"\mathrm{arg\,min}")
    tex = tex.replace(r"\operatorname*{arg\,max}", r"\mathrm{arg\,max}")
    tex = re.sub(r"\\begin\{subfigure\}(\[[^\]]*\])?\{[^}]*\}", r"\\begin{minipage}{0.48\\textwidth}", tex)
    tex = tex.replace(r"\end{subfigure}", r"\end{minipage}")
    return tex


def flatten_tex() -> str:
    raw = (BOOK / "book.tex").read_text(encoding="utf-8")
    raw = expand_inputs(raw, BOOK, skip={"preamble", "preamble.tex"})
    start = raw.find(r"\begin{document}")
    end = raw.rfind(r"\end{document}")
    if start < 0 or end < 0:
        raise SystemExit("book.tex missing document environment")
    body = raw[start + len(r"\begin{document}") : end]
    body = strip_comments(body)
    keys = extract_cite_keys(body)
    labels, eq_ids = collect_labels(body)
    body = replace_cites(body, keys)
    body = replace_refs(body, labels, eq_ids)
    body = convert_bibliography(body)
    body = convert_tabularx(body)
    body = strip_layout(body)
    return PANDOC_PREAMBLE + "\n" + body.strip() + "\n\\end{document}\n"


def set_rfonts(r_pr, east: str, ascii_font: str = TIMES) -> None:
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), ascii_font)
    r_fonts.set(qn("w:hAnsi"), ascii_font)
    r_fonts.set(qn("w:eastAsia"), east)
    r_fonts.set(qn("w:cs"), ascii_font)
    kern = r_pr.find(qn("w:kern"))
    if kern is None:
        kern = OxmlElement("w:kern")
        r_pr.append(kern)
    kern.set(qn("w:val"), "32")
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "zh-CN")


def style_font(style, east: str, size_pt: float, bold: bool = False) -> None:
    style.font.name = TIMES
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    r_pr = style.element.get_or_add_rPr()
    set_rfonts(r_pr, east)
    sz = r_pr.find(qn("w:sz"))
    if sz is None:
        sz = OxmlElement("w:sz")
        r_pr.append(sz)
    sz.set(qn("w:val"), str(int(size_pt * 2)))
    sz_cs = r_pr.find(qn("w:szCs"))
    if sz_cs is None:
        sz_cs = OxmlElement("w:szCs")
        r_pr.append(sz_cs)
    sz_cs.set(qn("w:val"), str(int(size_pt * 2)))
    if bold:
        b = r_pr.find(qn("w:b"))
        if b is None:
            b = OxmlElement("w:b")
            r_pr.append(b)


def set_style_spacing(style, before: float, after: float, line: float, indent_chars: int | None) -> None:
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    p_pr = style.element.get_or_add_pPr()
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:line"), str(int(line * 240)))
    spacing.set(qn("w:lineRule"), "auto")
    ind = p_pr.find(qn("w:ind"))
    if indent_chars is None:
        if ind is not None:
            p_pr.remove(ind)
        pf.first_line_indent = Pt(0)
        return
    if ind is None:
        ind = OxmlElement("w:ind")
        p_pr.append(ind)
    ind.set(qn("w:firstLineChars"), str(indent_chars))
    ind.set(qn("w:firstLine"), str(int(indent_chars * 2.4)))


def ensure_style(doc: Document, name: str) -> object:
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def set_doc_defaults(doc: Document) -> None:
    styles_el = doc.styles.element
    defaults = styles_el.find(qn("w:docDefaults"))
    if defaults is None:
        defaults = OxmlElement("w:docDefaults")
        styles_el.insert(0, defaults)
    r_default = defaults.find(qn("w:rPrDefault"))
    if r_default is None:
        r_default = OxmlElement("w:rPrDefault")
        defaults.append(r_default)
    r_pr = r_default.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = OxmlElement("w:rPr")
        r_default.append(r_pr)
    set_rfonts(r_pr, SONG)
    for tag, val in (("w:sz", "24"), ("w:szCs", "24")):
        node = r_pr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            r_pr.append(node)
        node.set(qn("w:val"), val)
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "zh-CN")
    p_default = defaults.find(qn("w:pPrDefault"))
    if p_default is None:
        p_default = OxmlElement("w:pPrDefault")
        defaults.append(p_default)
    p_pr = p_default.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = OxmlElement("w:pPr")
        p_default.append(p_pr)
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:line"), str(BODY_LINE))
    spacing.set(qn("w:lineRule"), "auto")
    ind = p_pr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        p_pr.append(ind)
    ind.set(qn("w:firstLineChars"), "200")
    ind.set(qn("w:firstLine"), "480")


def build_reference_doc(path: Path) -> None:
    doc = Document()
    set_doc_defaults(doc)
    section = doc.sections[0]
    section.page_width = A4_W
    section.page_height = A4_H
    section.orientation = WD_ORIENT.PORTRAIT
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)
    section.top_margin = Mm(24.5)
    section.bottom_margin = Mm(24.5)
    section.header_distance = Mm(12)
    section.footer_distance = Mm(12)

    normal = doc.styles["Normal"]
    style_font(normal, SONG, 12)
    set_style_spacing(normal, 0, 0, 1.5, 200)

    heading_specs = {
        "Heading 1": (HEI, 16, True, 18, 14, True),
        "Heading 2": (HEI, 14, True, 14, 6, False),
        "Heading 3": (HEI, 12, True, 10, 4, False),
        "Title": (HEI, 16, True, 0, 6, True),
        "Author": (SONG, 12, False, 0, 12, True),
    }
    for name, (east, size, bold, before, after, center) in heading_specs.items():
        style = ensure_style(doc, name)
        style_font(style, east, size, bold)
        set_style_spacing(style, before, after, 1.15, 0)
        if center:
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    body_names = [
        "First Paragraph",
        "Body Text",
        "Compact",
        "Block Text",
        "Definition",
        "Definition Term",
        "Source Code",
    ]
    for name in body_names:
        style = ensure_style(doc, name)
        style_font(style, SONG, 12, bold=(name == "Definition Term"))
        indent = 0 if name in {"Compact", "Definition Term"} else 200
        set_style_spacing(style, 0, 0, 1.5, indent)

    for name in ("Image Caption", "Table Caption", "Captioned Figure"):
        style = ensure_style(doc, name)
        style_font(style, SONG, 10.5)
        set_style_spacing(style, 4, 6, 1.2, 0)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(path)


def disable_auto_compress(doc: Document) -> None:
    settings = doc.settings.element
    for tag in ("w:doNotAutoCompressPictures", "w:displayBackgroundShape"):
        if settings.find(qn(tag)) is None:
            settings.append(OxmlElement(tag))
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    spacing = settings.find(qn("w:characterSpacingControl"))
    if spacing is None:
        spacing = OxmlElement("w:characterSpacingControl")
        settings.append(spacing)
    spacing.set(qn("w:val"), "compressPunctuation")


def configure_hyperlink_style(doc: Document) -> None:
    try:
        style = doc.styles["Hyperlink"]
    except KeyError:
        style = doc.styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
    r_pr = style.element.get_or_add_rPr()
    set_rfonts(r_pr, SONG)
    color = r_pr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        r_pr.append(color)
    color.set(qn("w:val"), "000000")
    underline = r_pr.find(qn("w:u"))
    if underline is None:
        underline = OxmlElement("w:u")
        r_pr.append(underline)
    underline.set(qn("w:val"), "none")


def run_font(run, east: str, size_pt: float | None = None, bold: bool | None = None) -> None:
    run.font.name = TIMES
    r_pr = run._element.get_or_add_rPr()
    set_rfonts(r_pr, east)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold


def prepend_run(paragraph, text: str, east: str, size_pt: float, bold: bool = False) -> None:
    r_el = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    set_rfonts(r_pr, east)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size_pt * 2)))
    r_pr.append(sz)
    if bold:
        r_pr.append(OxmlElement("w:b"))
    r_el.append(r_pr)
    t_el = OxmlElement("w:t")
    t_el.set(qn("xml:space"), "preserve")
    t_el.text = text
    r_el.append(t_el)
    p_el = paragraph._p
    p_pr = p_el.find(qn("w:pPr"))
    if p_pr is not None:
        p_pr.addnext(r_el)
    else:
        p_el.insert(0, r_el)


def clear_container(part) -> None:
    for paragraph in part.paragraphs:
        paragraph.clear()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.first_line_indent = Pt(0)


def add_bottom_border(paragraph, sz: str = "6") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "000000")


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    run._r.append(begin)
    instr = paragraph.add_run()._r
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = " PAGE "
    instr.append(text)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    paragraph.add_run()._r.append(sep)
    result = paragraph.add_run("1")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(end)
    for run in paragraph.runs:
        run_font(run, SONG, 10.5)


def inline_to_anchor(drawing, doc_pr_id: int, name: str) -> None:
    inline = drawing.find(qn("wp:inline"))
    if inline is None:
        return
    extent = inline.find(qn("wp:extent"))
    cx = extent.get("cx")
    cy = extent.get("cy")
    doc_pr = inline.find(qn("wp:docPr"))
    cnv = inline.find(qn("wp:cNvGraphicFramePr"))
    graphic = inline.find(qn("a:graphic"))
    doc_pr.set("id", str(doc_pr_id))
    doc_pr.set("name", name)
    anchor = OxmlElement("wp:anchor")
    for key, val in (
        ("distT", "0"),
        ("distB", "0"),
        ("distL", "0"),
        ("distR", "0"),
        ("simplePos", "0"),
        ("relativeHeight", "0"),
        ("behindDoc", "1"),
        ("locked", "1"),
        ("layoutInCell", "1"),
        ("allowOverlap", "1"),
    ):
        anchor.set(key, val)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    pos_h = OxmlElement("wp:positionH")
    pos_h.set("relativeFrom", "page")
    off_h = OxmlElement("wp:posOffset")
    off_h.text = "0"
    pos_h.append(off_h)
    pos_v = OxmlElement("wp:positionV")
    pos_v.set("relativeFrom", "page")
    off_v = OxmlElement("wp:posOffset")
    off_v.text = "0"
    pos_v.append(off_v)
    new_extent = OxmlElement("wp:extent")
    new_extent.set("cx", cx)
    new_extent.set("cy", cy)
    effect = OxmlElement("wp:effectExtent")
    for key in ("l", "t", "r", "b"):
        effect.set(key, "0")
    wrap = OxmlElement("wp:wrapNone")
    for child in (simple, pos_h, pos_v, new_extent, effect, wrap):
        anchor.append(child)
    if doc_pr is not None:
        anchor.append(doc_pr)
    if cnv is not None:
        anchor.append(cnv)
    if graphic is not None:
        anchor.append(graphic)
    drawing.remove(inline)
    drawing.append(anchor)


def make_watermark(dest: Path) -> None:
    image = Image.open(BG).convert("RGBA")
    image = image.resize((1240, 1754), Image.Resampling.LANCZOS)
    red, green, blue, alpha = image.split()
    alpha = alpha.point(lambda value: int(value * 0.13))
    Image.merge("RGBA", (red, green, blue, alpha)).save(dest, "PNG")


def add_watermark(section, image_path: Path) -> None:
    header = section.header
    paragraph = header.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.first_line_indent = Pt(0)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=A4_W, height=A4_H)
    drawing = run._r.find(qn("w:drawing"))
    if drawing is not None:
        inline_to_anchor(drawing, 9001, "watermark")


def set_page_start(section, number: int) -> None:
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(number))


def configure_section(section, full_bleed: bool) -> None:
    section.page_width = A4_W
    section.page_height = A4_H
    section.orientation = WD_ORIENT.PORTRAIT
    if full_bleed:
        section.left_margin = Mm(0)
        section.right_margin = Mm(0)
        section.top_margin = Mm(0)
        section.bottom_margin = Mm(0)
        section.header_distance = Mm(0)
        section.footer_distance = Mm(0)
    else:
        section.left_margin = Mm(25)
        section.right_margin = Mm(25)
        section.top_margin = Mm(24.5)
        section.bottom_margin = Mm(24.5)
        section.header_distance = Mm(12)
        section.footer_distance = Mm(12)
    start = section._sectPr.find(qn("w:type"))
    if start is None:
        start = OxmlElement("w:type")
        section._sectPr.append(start)
    start.set(qn("w:val"), "nextPage")
    grid = section._sectPr.find(qn("w:docGrid"))
    if grid is None:
        grid = OxmlElement("w:docGrid")
        section._sectPr.append(grid)
    grid.set(qn("w:type"), "none")


def make_fullbleed_sectpr() -> OxmlElement:
    sect_pr = OxmlElement("w:sectPr")
    pg_sz = OxmlElement("w:pgSz")
    pg_sz.set(qn("w:w"), "11906")
    pg_sz.set(qn("w:h"), "16838")
    sect_pr.append(pg_sz)
    pg_mar = OxmlElement("w:pgMar")
    for attr in ("w:top", "w:right", "w:bottom", "w:left", "w:header", "w:footer", "w:gutter"):
        pg_mar.set(qn(attr), "0")
    sect_pr.append(pg_mar)
    typ = OxmlElement("w:type")
    typ.set(qn("w:val"), "nextPage")
    sect_pr.append(typ)
    return sect_pr


def add_full_page_picture(paragraph, image_path: Path) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.first_line_indent = Pt(0)
    p_pr = paragraph._p.get_or_add_pPr()
    ind = p_pr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        p_pr.append(ind)
    ind.set(qn("w:firstLine"), "0")
    ind.set(qn("w:firstLineChars"), "0")
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=A4_W, height=A4_H)


def insert_cover(doc: Document, image_path: Path) -> None:
    first = doc.paragraphs[0]
    cover = first.insert_paragraph_before()
    add_full_page_picture(cover, image_path)
    cover._p.get_or_add_pPr().append(make_fullbleed_sectpr())


def setup_headers(doc: Document, watermark: Path) -> None:
    cover, body, back = doc.sections[0], doc.sections[1], doc.sections[2]
    configure_section(cover, True)
    configure_section(body, False)
    configure_section(back, True)
    set_page_start(body, 1)

    for section in (cover, back):
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        clear_container(section.header)
        clear_container(section.footer)

    body.header.is_linked_to_previous = False
    body.footer.is_linked_to_previous = False
    clear_container(body.header)
    clear_container(body.footer)

    header_p = body.header.paragraphs[0] if body.header.paragraphs else body.header.add_paragraph()
    header_p.paragraph_format.first_line_indent = Pt(0)
    header_p.paragraph_format.space_before = Pt(0)
    header_p.paragraph_format.space_after = Pt(4)
    header_p.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT)
    left = header_p.add_run(HEADER_LEFT)
    run_font(left, SONG, 9)
    header_p.add_run("\t")
    right = header_p.add_run(HEADER_RIGHT)
    run_font(right, SONG, 9)
    add_bottom_border(header_p, "8")
    add_watermark(body, watermark)

    footer_p = body.footer.paragraphs[0] if body.footer.paragraphs else body.footer.add_paragraph()
    footer_p.paragraph_format.first_line_indent = Pt(0)
    add_page_field(footer_p)


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def set_bool_prop(p_pr, tag: str, enabled: bool = True) -> None:
    node = p_pr.find(qn(tag))
    if not enabled:
        if node is not None:
            p_pr.remove(node)
        return
    if node is None:
        node = OxmlElement(tag)
        p_pr.append(node)


def set_paragraph_layout(
    paragraph,
    *,
    justify: bool | None = None,
    first_line_chars: int | None = None,
    left_twips: int | None = None,
    hanging_twips: int | None = None,
    page_break: bool = False,
    keep_next: bool = False,
    before: int | None = None,
    after: int | None = None,
    line: int | None = 331,
) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    set_bool_prop(p_pr, "w:autoSpaceDE", False)
    auto_de = p_pr.find(qn("w:autoSpaceDE"))
    if auto_de is None:
        auto_de = OxmlElement("w:autoSpaceDE")
        p_pr.append(auto_de)
    auto_de.set(qn("w:val"), "0")
    auto_dn = p_pr.find(qn("w:autoSpaceDN"))
    if auto_dn is None:
        auto_dn = OxmlElement("w:autoSpaceDN")
        p_pr.append(auto_dn)
    auto_dn.set(qn("w:val"), "0")
    set_bool_prop(p_pr, "w:kinsoku", True)
    set_bool_prop(p_pr, "w:wordWrap", True)
    set_bool_prop(p_pr, "w:overflowPunct", True)
    set_bool_prop(p_pr, "w:widowControl", True)
    snap = p_pr.find(qn("w:snapToGrid"))
    if snap is None:
        snap = OxmlElement("w:snapToGrid")
        p_pr.append(snap)
    snap.set(qn("w:val"), "0")
    if justify is True:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif justify is False:
        pass
    if page_break:
        set_bool_prop(p_pr, "w:pageBreakBefore", True)
    if keep_next:
        set_bool_prop(p_pr, "w:keepNext", True)
        set_bool_prop(p_pr, "w:keepLines", True)
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    if line is not None:
        spacing.set(qn("w:line"), str(line))
        spacing.set(qn("w:lineRule"), "auto")
    if before is not None:
        spacing.set(qn("w:before"), str(before))
    if after is not None:
        spacing.set(qn("w:after"), str(after))
    ind = p_pr.find(qn("w:ind"))
    if ind is None and (
        first_line_chars is not None or left_twips is not None or hanging_twips is not None
    ):
        ind = OxmlElement("w:ind")
        p_pr.append(ind)
    if ind is not None:
        if first_line_chars is not None:
            ind.set(qn("w:firstLineChars"), str(first_line_chars))
            ind.set(qn("w:firstLine"), str(int(first_line_chars * 2.4)))
            if qn("w:hanging") in ind.attrib:
                del ind.attrib[qn("w:hanging")]
        if left_twips is not None:
            ind.set(qn("w:left"), str(left_twips))
        if hanging_twips is not None:
            ind.set(qn("w:left"), str(hanging_twips))
            ind.set(qn("w:hanging"), str(hanging_twips))
            if qn("w:firstLine") in ind.attrib:
                del ind.attrib[qn("w:firstLine")]
            if qn("w:firstLineChars") in ind.attrib:
                del ind.attrib[qn("w:firstLineChars")]


def spaced_cn(text: str) -> str:
    if len(text) == 2 and all("\u4e00" <= ch <= "\u9fff" for ch in text):
        return text[0] + "　" + text[1]
    return text


def heading_toc_text(paragraph) -> str:
    chunks = []
    for run in paragraph.runs:
        raw = (run.text or "").replace("\n", "").replace("\r", "").replace("\t", "")
        chunks.append(raw)
        if run._r.find(qn("w:br")) is not None:
            chunks.append("　")
    text = re.sub(r"　+", "　", "".join(chunks)).strip()
    return text or re.sub(r"\s+", "　", paragraph.text).strip()


def write_chapter_heading(paragraph, kicker: str, title: str) -> None:
    paragraph.clear()
    kicker_run = paragraph.add_run(kicker)
    run_font(kicker_run, HEI, 14, True)
    kicker_run._r.append(OxmlElement("w:br"))
    title_run = paragraph.add_run(spaced_cn(title))
    run_font(title_run, HEI, 18, True)
    r_pr = title_run._element.get_or_add_rPr()
    spacing = r_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        r_pr.append(spacing)
    spacing.set(qn("w:val"), "60")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    p_el = paragraph._p
    p_pr = p_el.find(qn("w:pPr"))
    if p_pr is not None:
        p_pr.addnext(start)
    else:
        p_el.insert(0, start)
    p_el.append(end)


def add_heading_bookmarks(doc: Document) -> list[tuple[int, str, str]]:
    entries: list[tuple[int, str, str]] = []
    bookmark_id = 1
    for paragraph in doc.paragraphs:
        name = paragraph.style.name if paragraph.style is not None else ""
        text = paragraph.text.strip()
        if not text:
            continue
        level = {"Heading 1": 1, "Heading 2": 2, "Heading 3": 3}.get(name)
        if level is None:
            continue
        mark = f"_Toc{bookmark_id:04d}"
        add_bookmark(paragraph, mark, bookmark_id)
        entries.append((level, heading_toc_text(paragraph) or text, mark))
        bookmark_id += 1
    return entries


def append_field_runs(parent, instruction: str, east: str, size_pt: float, bold: bool) -> None:
    def add_run(build) -> None:
        run = OxmlElement("w:r")
        r_pr = OxmlElement("w:rPr")
        set_rfonts(r_pr, east)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size_pt * 2)))
        r_pr.append(sz)
        if bold:
            r_pr.append(OxmlElement("w:b"))
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "000000")
        r_pr.append(color)
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "none")
        r_pr.append(underline)
        run.append(r_pr)
        build(run)
        parent.append(run)

    def begin(run) -> None:
        node = OxmlElement("w:fldChar")
        node.set(qn("w:fldCharType"), "begin")
        run.append(node)

    def instr(run) -> None:
        node = OxmlElement("w:instrText")
        node.set(qn("xml:space"), "preserve")
        node.text = instruction
        run.append(node)

    def separate(run) -> None:
        node = OxmlElement("w:fldChar")
        node.set(qn("w:fldCharType"), "separate")
        run.append(node)

    def placeholder(run) -> None:
        text = OxmlElement("w:t")
        text.set(qn("xml:space"), "preserve")
        text.text = " "
        run.append(text)

    def end(run) -> None:
        node = OxmlElement("w:fldChar")
        node.set(qn("w:fldCharType"), "end")
        run.append(node)

    add_run(begin)
    add_run(instr)
    add_run(separate)
    add_run(placeholder)
    add_run(end)


def number_headings(doc: Document) -> None:
    body_map = {title: index + 1 for index, title in enumerate(BODY_CHAPTERS)}
    app_map = {title: chr(ord("A") + index) for index, title in enumerate(APPENDIX_CHAPTERS)}
    current: int | str | None = None
    sec = 0
    sub = 0
    stray = []
    for paragraph in doc.paragraphs:
        name = paragraph.style.name if paragraph.style is not None else ""
        title = paragraph.text.strip()
        if name in {"Title", "Author"}:
            stray.append(paragraph)
            continue
        if name == "Heading 1":
            if title in UNNUMBERED:
                current = None
                sec = 0
                continue
            if title in body_map:
                current = body_map[title]
                sec = 0
                sub = 0
                write_chapter_heading(paragraph, f"第{current}章", title)
                continue
            if title in app_map:
                current = app_map[title]
                sec = 0
                sub = 0
                write_chapter_heading(paragraph, f"附录 {current}", title)
                continue
        if name == "Heading 2" and current is not None:
            sec += 1
            sub = 0
            paragraph.text = f"{current}.{sec}　{title}"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif name == "Heading 3" and current is not None:
            sub += 1
            paragraph.text = f"{current}.{sec}.{sub}　{title}"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for paragraph in stray:
        delete_paragraph(paragraph)


def number_captions(doc: Document) -> None:
    current: int | str | None = None
    fig = 0
    tab = 0
    for paragraph in doc.paragraphs:
        name = paragraph.style.name if paragraph.style is not None else ""
        text = paragraph.text.strip()
        if name == "Heading 1":
            if text.startswith("第") and "章" in text:
                try:
                    current = int(text[1 : text.index("章")])
                except ValueError:
                    current = None
                fig = tab = 0
            elif text.startswith("附录"):
                match = re.search(r"附录\s*([A-Z])", text)
                current = match.group(1) if match else None
                fig = tab = 0
            else:
                current = None
                fig = tab = 0
            continue
        if current is None:
            continue
        if name == "Image Caption" and text and not text.startswith("图"):
            fig += 1
            prepend_run(paragraph, f"图 {current}.{fig}  ", HEI, 10.5, True)
        elif name == "Table Caption" and text and not text.startswith("表"):
            tab += 1
            prepend_run(paragraph, f"表 {current}.{tab}  ", HEI, 10.5, True)


def toc_run_props(east: str, size_pt: float, bold: bool) -> OxmlElement:
    r_pr = OxmlElement("w:rPr")
    set_rfonts(r_pr, east)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size_pt * 2)))
    r_pr.append(sz)
    sz_cs = OxmlElement("w:szCs")
    sz_cs.set(qn("w:val"), str(int(size_pt * 2)))
    r_pr.append(sz_cs)
    if bold:
        r_pr.append(OxmlElement("w:b"))
        b_cs = OxmlElement("w:bCs")
        r_pr.append(b_cs)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "none")
    r_pr.append(underline)
    return r_pr


def insert_toc(doc: Document, entries: list[tuple[int, str, str]]) -> None:
    anchor = None
    for paragraph in doc.paragraphs:
        if paragraph.style is None:
            continue
        if paragraph.style.name == "Heading 1" and paragraph.text.startswith("第1章"):
            anchor = paragraph
            break
    if anchor is None or not entries:
        return
    first_item = None
    for level, text, mark in entries:
        item = anchor.insert_paragraph_before()
        east = HEI if level == 1 else SONG
        size = {1: 12, 2: 11, 3: 10.5}[level]
        bold = level == 1
        left = {1: 0, 2: 360, 3: 720}[level]
        before = {1: 160, 2: 40, 3: 20}[level]
        after = {1: 40, 2: 20, 3: 10}[level]
        set_paragraph_layout(
            item,
            justify=False,
            first_line_chars=0,
            left_twips=left,
            before=before,
            after=after,
            line=TOC_LINE,
        )
        item.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_pr = item._p.get_or_add_pPr()
        tabs = p_pr.find(qn("w:tabs"))
        if tabs is None:
            tabs = OxmlElement("w:tabs")
            p_pr.append(tabs)
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:leader"), "dot")
        tab.set(qn("w:pos"), TOC_TAB_TWIPS)
        tabs.append(tab)
        ind = p_pr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            p_pr.append(ind)
        ind.set(qn("w:right"), "200")

        hyper = OxmlElement("w:hyperlink")
        hyper.set(qn("w:anchor"), mark)
        hyper.set(qn("w:history"), "1")
        title_run = OxmlElement("w:r")
        title_run.append(toc_run_props(east, size, bold))
        title_text = OxmlElement("w:t")
        title_text.set(qn("xml:space"), "preserve")
        title_text.text = text
        title_run.append(title_text)
        hyper.append(title_run)
        tab_run = OxmlElement("w:r")
        tab_run.append(toc_run_props(east, size, bold))
        tab_run.append(OxmlElement("w:tab"))
        hyper.append(tab_run)
        append_field_runs(hyper, f" PAGEREF {mark} \\h ", east, size, bold)
        item._p.append(hyper)
        if first_item is None:
            first_item = item
    rule = first_item.insert_paragraph_before()
    set_paragraph_layout(rule, justify=False, first_line_chars=0, before=0, after=80, line=240)
    add_bottom_border(rule, "6")
    title = rule.insert_paragraph_before("目　录")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run_font(run, HEI, 18, True)
        r_pr = run._element.get_or_add_rPr()
        spacing = r_pr.find(qn("w:spacing"))
        if spacing is None:
            spacing = OxmlElement("w:spacing")
            r_pr.append(spacing)
        spacing.set(qn("w:val"), "160")
    set_paragraph_layout(
        title,
        justify=False,
        first_line_chars=0,
        page_break=True,
        keep_next=True,
        before=360,
        after=80,
        line=HEAD_LINE,
    )


def restyle_runs(doc: Document) -> None:
    heading_sizes = {"Heading 1": (HEI, 16, True), "Heading 2": (HEI, 14, True), "Heading 3": (HEI, 12, True)}
    caption_names = {"Image Caption", "Table Caption", "Captioned Figure"}
    for paragraph in doc.paragraphs:
        name = paragraph.style.name if paragraph.style is not None else "Normal"
        if name == "Heading 1":
            runs = paragraph.runs
            if len(runs) >= 2 and runs[0]._r.find(qn("w:br")) is not None:
                run_font(runs[0], HEI, 14, True)
                for run in runs[1:]:
                    run_font(run, HEI, 18, True)
            else:
                for run in runs:
                    run_font(run, HEI, 16, True)
            continue
        if name in heading_sizes:
            east, size, bold = heading_sizes[name]
            for run in paragraph.runs:
                run_font(run, east, size, bold)
            continue
        if name in caption_names:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for index, run in enumerate(paragraph.runs):
                if index == 0:
                    run_font(run, HEI, 10.5, True)
                else:
                    run_font(run, KAI, 10.5, False)
            continue
        if name in {"Title", "Author"}:
            continue
        for run in paragraph.runs:
            run_font(run, SONG, 12)
    style_tables(doc)


def style_tables(doc: Document) -> None:
    for table in doc.tables:
        tbl = table._tbl
        tbl_pr = tbl.tblPr
        if tbl_pr is None:
            continue
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:type"), "pct")
        tbl_w.set(qn("w:w"), "5000")
        jc = tbl_pr.find(qn("w:jc"))
        if jc is None:
            jc = OxmlElement("w:jc")
            tbl_pr.append(jc)
        jc.set(qn("w:val"), "center")
        borders = tbl_pr.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            tbl_pr.append(borders)
        for child in list(borders):
            borders.remove(child)
        for edge, size, val in (
            ("top", "12", "single"),
            ("left", "0", "nil"),
            ("bottom", "12", "single"),
            ("right", "0", "nil"),
            ("insideH", "4", "single"),
            ("insideV", "0", "nil"),
        ):
            node = OxmlElement(f"w:{edge}")
            node.set(qn("w:val"), val)
            node.set(qn("w:sz"), size)
            node.set(qn("w:space"), "0")
            node.set(qn("w:color"), "000000")
            borders.append(node)
        cell_mar = tbl_pr.find(qn("w:tblCellMar"))
        if cell_mar is None:
            cell_mar = OxmlElement("w:tblCellMar")
            tbl_pr.append(cell_mar)
        for edge, value in (("top", "40"), ("left", "80"), ("bottom", "40"), ("right", "80")):
            node = cell_mar.find(qn(f"w:{edge}"))
            if node is None:
                node = OxmlElement(f"w:{edge}")
                cell_mar.append(node)
            node.set(qn("w:w"), value)
            node.set(qn("w:type"), "dxa")
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                tc = cell._tc
                tc_pr = tc.get_or_add_tcPr()
                v_align = tc_pr.find(qn("w:vAlign"))
                if v_align is None:
                    v_align = OxmlElement("w:vAlign")
                    tc_pr.append(v_align)
                v_align.set(qn("w:val"), "center")
                for paragraph in cell.paragraphs:
                    set_paragraph_layout(
                        paragraph,
                        justify=False,
                        first_line_chars=0,
                        before=40,
                        after=40,
                        line=276,
                    )
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run_font(run, HEI if row_index == 0 else SONG, 10.5, row_index == 0)


def polish_layout(doc: Document) -> None:
    in_refs = False
    for paragraph in doc.paragraphs:
        name = paragraph.style.name if paragraph.style is not None else "Normal"
        text = paragraph.text.strip()
        xml = paragraph._p.xml
        if name == "Heading 1":
            in_refs = text.startswith("参考文献")
            page_break = text not in {"摘要", "Abstract"}
            two_line = any(run._r.find(qn("w:br")) is not None for run in paragraph.runs)
            set_paragraph_layout(
                paragraph,
                justify=False,
                first_line_chars=0,
                page_break=page_break,
                keep_next=True,
                before=480 if two_line else 280,
                after=280 if two_line else 200,
                line=320 if two_line else HEAD_LINE,
            )
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if name == "Heading 2":
            set_paragraph_layout(
                paragraph,
                justify=False,
                first_line_chars=0,
                keep_next=True,
                before=260,
                after=80,
                line=HEAD_LINE,
            )
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            continue
        if name == "Heading 3":
            set_paragraph_layout(
                paragraph,
                justify=False,
                first_line_chars=0,
                keep_next=True,
                before=180,
                after=60,
                line=HEAD_LINE,
            )
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            continue
        if name in {"Image Caption", "Table Caption"}:
            set_paragraph_layout(
                paragraph,
                justify=False,
                first_line_chars=0,
                keep_next=(name == "Table Caption"),
                before=80 if name == "Image Caption" else 160,
                after=160 if name == "Image Caption" else 80,
                line=HEAD_LINE,
            )
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if name == "Captioned Figure" or "<w:drawing" in xml:
            set_paragraph_layout(
                paragraph,
                justify=False,
                first_line_chars=0,
                keep_next=True,
                before=180,
                after=40,
                line=240,
            )
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if "m:oMath" in xml and len(text) < 12:
            set_paragraph_layout(
                paragraph,
                justify=False,
                first_line_chars=0,
                before=100,
                after=100,
                line=HEAD_LINE,
            )
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if in_refs and name != "Heading 1":
            set_paragraph_layout(
                paragraph,
                justify=True,
                hanging_twips=420,
                before=40,
                after=40,
                line=HEAD_LINE,
            )
            for run in paragraph.runs:
                run_font(run, TIMES, 10.5)
            continue
        if name in {"Compact", "List Number", "List Bullet", "List Paragraph"}:
            set_paragraph_layout(
                paragraph,
                justify=True,
                first_line_chars=0,
                left_twips=420,
                before=40,
                after=40,
                line=BODY_LINE,
            )
            continue
        set_paragraph_layout(
            paragraph,
            justify=True,
            first_line_chars=200,
            before=0,
            after=0,
            line=BODY_LINE,
        )


def resize_body_figures(doc: Document) -> None:
    max_emu = int(BODY_FIG_MAX)
    page_emu = int(A4_W)
    for shape in doc.inline_shapes:
        width = int(shape.width)
        if width >= page_emu - 20000:
            continue
        if width > max_emu:
            height = int(shape.height)
            shape.width = max_emu
            shape.height = int(height * (max_emu / width))


def center_figure_paragraphs(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        name = paragraph.style.name if paragraph.style is not None else ""
        if name in {"Captioned Figure", "Image Caption"}:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Pt(0)
            continue
        drawing = paragraph._p.find(
            ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing"
        )
        if drawing is not None:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Pt(0)


def find_pandoc() -> str:
    found = shutil.which("pandoc")
    if found:
        return found
    fallback = Path(r"D:\Anaconda3\Scripts\pandoc.exe")
    if fallback.exists():
        return str(fallback)
    raise SystemExit("pandoc not found")


def run_pandoc(tex_path: Path, docx_path: Path, reference: Path, media: Path) -> None:
    resource = ";".join(
        [
            ".",
            "figures",
            str(BOOK / "figures"),
            str(ROOT / "results" / "figures" / "paper"),
            str(ROOT / "results" / "figures" / "study"),
            str(ROOT / "results" / "figures"),
        ]
    )
    cmd = [
        find_pandoc(),
        str(tex_path),
        "-o",
        str(docx_path),
        "--from",
        "latex",
        "--to",
        "docx",
        "--wrap=none",
        "--top-level-division=chapter",
        "--reference-doc",
        str(reference),
        "--resource-path",
        resource,
        "--extract-media",
        str(media),
        "--metadata",
        "lang=zh-CN",
    ]
    result = subprocess.run(cmd, cwd=str(BOOK), capture_output=True, text=True, encoding="utf-8", errors="replace")
    if result.returncode != 0:
        print(result.stdout)
        print(result.stderr)
        raise SystemExit(f"pandoc failed: {result.returncode}")
    if result.stderr:
        for line in result.stderr.splitlines():
            if "WARNING" in line or "Could not convert" in line:
                print(line)


def extract_toc_pages(xml: str) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for match in re.finditer(r"PAGEREF\s+(_Toc\d+)", xml):
        mark = match.group(1)
        sep = xml.find('w:fldCharType="separate"', match.end())
        if sep < 0 or sep - match.end() > 1200:
            continue
        t_open = xml.find("<w:t", sep)
        if t_open < 0 or t_open - sep > 800:
            continue
        t_gt = xml.find(">", t_open)
        t_end = xml.find("</w:t>", t_gt)
        if t_gt < 0 or t_end < 0:
            continue
        value = xml[t_gt + 1 : t_end].strip()
        if re.fullmatch(r"\d+", value):
            mapping[mark] = value
    return mapping


def patch_toc_pages(xml: str, mapping: dict[str, str]) -> str:
    out = xml
    for mark, number in mapping.items():
        needle = f"PAGEREF {mark} \\h "
        start = 0
        while True:
            index = out.find(needle, start)
            if index < 0:
                break
            sep = out.find('w:fldCharType="separate"', index)
            if sep < 0:
                break
            t_open = out.find("<w:t", sep)
            if t_open < 0 or t_open - sep > 800:
                break
            t_gt = out.find(">", t_open)
            t_end = out.find("</w:t>", t_gt)
            if t_gt < 0 or t_end < 0:
                break
            out = out[: t_gt + 1] + number + out[t_end:]
            start = t_gt + 1 + len(number)
    return out


def rewrite_zip_entry(docx_path: Path, name: str, xml: str) -> None:
    tmp_path = docx_path.with_name(docx_path.name + ".tmp")
    with ZipFile(docx_path, "r") as source, ZipFile(tmp_path, "w") as dest:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == name:
                data = xml.encode("utf-8")
            copied = ZipInfo(filename=info.filename, date_time=info.date_time)
            copied.compress_type = info.compress_type or ZIP_DEFLATED
            dest.writestr(copied, data)
    tmp_path.replace(docx_path)


def update_fields_with_word(docx_path: Path) -> None:
    import pythoncom
    import win32com.client

    abs_path = str(docx_path.resolve())
    pythoncom.CoInitialize()
    word = None
    document = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        try:
            word.ScreenUpdating = False
        except Exception:
            pass
        document = word.Documents.Open(
            abs_path,
            ConfirmConversions=False,
            ReadOnly=False,
            AddToRecentFiles=False,
        )
        document.Repaginate()
        document.Fields.Update()
        for section in document.Sections:
            for header in section.Headers:
                header.Range.Fields.Update()
            for footer in section.Footers:
                footer.Range.Fields.Update()
        document.Repaginate()
        document.Fields.Update()
        document.Save()
        document.Close(False)
        document = None
    finally:
        if document is not None:
            try:
                document.Close(False)
            except Exception:
                pass
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()


def bake_toc_pages(docx_path: Path, work: Path) -> None:
    worded = work / "worded.docx"
    shutil.copy2(docx_path, worded)
    print("updating Word fields for TOC page numbers...")
    update_fields_with_word(worded)
    with ZipFile(worded) as archive:
        baked = archive.read("word/document.xml").decode("utf-8")
    mapping = extract_toc_pages(baked)
    with ZipFile(docx_path) as archive:
        original = archive.read("word/document.xml").decode("utf-8")
    marks = re.findall(r'w:anchor="(_Toc\d+)"', original)
    last = "1"
    for mark in marks:
        value = mapping.get(mark, "").strip()
        if re.fullmatch(r"\d+", value):
            last = value
        else:
            mapping[mark] = last
    if len(mapping) < 20:
        raise SystemExit(f"Word filled only {len(mapping)} TOC page numbers")
    patched = patch_toc_pages(original, mapping)
    filled = extract_toc_pages(patched)
    missing = [mark for mark in marks if mark not in filled]
    if missing:
        raise SystemExit(f"patched TOC missing pages: {missing[:8]}")
    rewrite_zip_entry(docx_path, "word/document.xml", patched)
    numbers = sorted(int(value) for value in filled.values())
    sample = ", ".join(f"{mark}={filled[mark]}" for mark in list(filled)[:4])
    print(f"TOC pages baked: {len(filled)} entries, {numbers[0]}–{numbers[-1]} ({sample})")


def verify_native(docx_path: Path) -> None:
    with ZipFile(docx_path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
        names = archive.namelist()
        media = [name for name in names if name.startswith("word/media/")]
    if "蒸烤一体" not in xml:
        raise SystemExit("docx has no body text")
    if xml.count("<w:t") < 400:
        raise SystemExit("docx still looks like page images")
    if xml.count("m:oMath") < 200:
        raise SystemExit("docx lost equations")
    if len(media) < 40:
        raise SystemExit("docx lost figures")
    size_mb = docx_path.stat().st_size / 1e6
    if size_mb > 90:
        raise SystemExit(f"docx too large: {size_mb:.1f} MB")
    if xml.count("w:hyperlink") < 20:
        raise SystemExit("docx TOC lost hyperlinks")
    if xml.count("w:bookmarkStart") < 20:
        raise SystemExit("docx lost heading bookmarks")
    if "目　录" not in xml:
        raise SystemExit("docx lost TOC title")
    pages = extract_toc_pages(xml)
    marks = re.findall(r'w:anchor="(_Toc\d+)"', xml)
    missing = [mark for mark in marks if mark not in pages]
    if missing:
        raise SystemExit(f"docx TOC page numbers empty: {missing[:8]}")
    print(
        "native docx ok:",
        f"w:t={xml.count('<w:t')}",
        f"omml={xml.count('m:oMath')}",
        f"media={len(media)}",
        f"hyper={xml.count('w:hyperlink')}",
        f"toc_pages={len(pages)}",
        f"size={size_mb:.1f}MB",
    )


def build() -> None:
    if not COVER.exists() or not BACK.exists():
        raise SystemExit("missing cover/back; run scripts/draw_covers.py")
    work = Path(tempfile.mkdtemp(prefix="acts-docx-"))
    try:
        tex_path = work / "book_flat.tex"
        ref_path = work / "reference.docx"
        body_path = work / "body.docx"
        media = work / "media"
        watermark = work / "watermark.png"
        tex_path.write_text(flatten_tex(), encoding="utf-8")
        build_reference_doc(ref_path)
        run_pandoc(tex_path, body_path, ref_path, media)
        make_watermark(watermark)

        doc = Document(str(body_path))
        disable_auto_compress(doc)
        configure_hyperlink_style(doc)
        number_headings(doc)
        toc_entries = add_heading_bookmarks(doc)
        number_captions(doc)
        restyle_runs(doc)
        polish_layout(doc)
        resize_body_figures(doc)
        center_figure_paragraphs(doc)
        insert_toc(doc, toc_entries)
        insert_cover(doc, COVER)
        back_section = doc.add_section(WD_SECTION.NEW_PAGE)
        configure_section(back_section, True)
        add_full_page_picture(doc.add_paragraph(), BACK)
        setup_headers(doc, watermark)

        SUBMIT_DIR.mkdir(exist_ok=True)
        if PDF.exists():
            SUBMIT_PDF.write_bytes(PDF.read_bytes())
        doc.save(SUBMIT_DOCX)
        bake_toc_pages(SUBMIT_DOCX, work)
        shutil.copy2(SUBMIT_DOCX, LOCAL_DOCX)
        verify_native(SUBMIT_DOCX)
        print("wrote", SUBMIT_DOCX)
        print("wrote", LOCAL_DOCX)
    finally:
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    build()
